"""Anvil workbook compiler and server-side binary scoring. No hosted services."""
import hashlib
import json
import random
import re
import secrets
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from core.scoring import normalize

MEDIA = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.mp3', '.wav', '.m4a', '.mp4', '.webm', '.ogg'}
IMAGES = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp'}
BRACKET = re.compile(r'\[([^\[\]]*)\]')


def _text(value):
    return '' if value is None else str(value).strip()


def _id(value):
    return hashlib.sha256(value.encode()).hexdigest()[:24]


def compile_source(root, manifest, asset):
    root = Path(root).resolve()
    source = manifest.get('source', 'course.xlsx')
    if not isinstance(source, str) or '\\' in source or ':' in source:
        raise ValueError('source: invalid workbook path')
    path = (root / source).resolve()
    if root not in path.parents or path.suffix.lower() != '.xlsx' or not path.is_file():
        raise ValueError('source: workbook must exist inside the course ZIP')
    if type(manifest.get('legacy', False)) is not bool:
        raise ValueError('legacy: expected boolean')
    legacy = manifest.get('legacy', False)
    try:
        with ZipFile(path) as archive:
            if sum(member.file_size for member in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError('Workbook expanded content exceeds 64 MiB')
        book = load_workbook(path, read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException) as error:
        raise ValueError('source: invalid XLSX workbook') from error
    sheet = book.active
    if sheet.max_row > 20000 or sheet.max_column > 256 or sheet.max_row * sheet.max_column > 1000000:
        book.close()
        raise ValueError('Workbook exceeds supported row/cell limits')
    raw_rows = list(sheet.iter_rows(values_only=True))
    book.close()
    if not raw_rows:
        raise ValueError('Workbook is empty')
    if not legacy and [str(v or '').lower().strip() for v in raw_rows[0][:4]] != ['difficulty','type','stimulus','prompt']:
        raise ValueError('A1:D1: expected difficulty, type, stimulus, prompt')
    rows = []
    for number, raw in enumerate(raw_rows[1:], 2):
        values = [_text(v) for v in raw]
        if not any(values):
            continue
        values += [''] * max(0, 7-len(values))
        if legacy:
            values = ['1', '', *values]
        rows.append((number, values))
    questions, items, warnings = [], [], []
    rng = random.Random(int(hashlib.sha256(path.read_bytes()).hexdigest(),16))

    def unit(raw, row, column):
        raw = _text(raw)
        cell = f'{get_column_letter(column-2 if legacy else column)}{row}'
        if not raw or raw.startswith('='):
            raise ValueError(f'{cell}: empty/formula content is not supported')
        if Path(raw).suffix.lower() in {'.avi','.mov','.mkv','.wmv','.flv','.svg','.html','.exe'}:
            raise ValueError(f'{cell}: unsupported media format; convert before upload')
        if Path(raw).suffix.lower() in MEDIA:
            try:
                return {'kind': 'media', 'media': asset(raw)}
            except (ValueError, OSError) as error:
                raise ValueError(f'{cell}: {error}') from error
        result = {'kind':'text', 'text':raw}
        # Source-coordinate audio is optional and never synthesized.
        narration = f'{cell}.mp3'
        if (root / narration).is_file():
            result['speechUrl'] = asset(narration)['path']
        return result

    def choices(raw, row, start=5):
        if len(raw) < 2 or len(raw) > 64 or len(set(raw)) != len(raw):
            raise ValueError(f'{get_column_letter(start-2 if legacy else start)}{row}: need 2–64 distinct choices')
        options = [{'id':secrets.token_hex(12), 'content':unit(value,row,start+j)} for j,value in enumerate(raw)]
        expected = options[0]['id']
        labels = {o['id']: raw[j] for j,o in enumerate(options)}
        rng.shuffle(options)
        return options, expected, labels

    def add(row, subtype, prompt, interaction, solution, suffix=''):
        ident = _id(f'anvil|{source}|{row}|{suffix}')
        name = prompt or f'Anvil {subtype} · {row}'
        items.append({'id':ident,'name':name,'values':{},'difficulty':int(difficulty or '1')})
        questions.append({'objectId':ident,'kind':f'anvil-{subtype}','prompt':prompt,'interaction':interaction,'solution':solution})

    header = None
    for index, (row, values) in enumerate(rows):
        difficulty, declared, stimulus, prompt = values[:4]
        raw = values[4:]
        while raw and not raw[-1]:
            raw.pop()
        try:
            if stimulus and not difficulty.isdigit():
                raise ValueError(f'A{row}: difficulty must be a nonnegative integer')
            if not stimulus and not declared and header:
                kind, shared, labels = header
                if kind == 2:
                    if not prompt or not raw or raw[0] not in labels:
                        raise ValueError(f'D{row}: matrix subquestion must name an exact header option')
                    content = unit(prompt,row,4)
                    add(row,2,content.get('text',''),dict(shared,question=content),{'mode':'choice','expected':labels[raw[0]],'labels':{v:k for k,v in labels.items()}})
                    continue
                if kind == 3:
                    options, expected, labels = choices(raw,row)
                    content = unit(prompt,row,4)
                    add(row,3,content.get('text',''),dict(shared,question=content,options=options),{'mode':'choice','expected':expected,'labels':labels})
                    continue
            header = None
            if legacy:
                if stimulus == '*':
                    kind = 8 if raw and all(Path(v.rsplit('*',1)[-1].strip()).suffix.lower() in MEDIA for v in raw) else 6
                elif '[' in stimulus:
                    kind = 5 if '[~' in stimulus else 4
                elif Path(stimulus).suffix.lower() in MEDIA:
                    next_row = rows[index+1][1] if index+1 < len(rows) else None
                    if raw and all(Path(v).suffix.lower() in IMAGES for v in raw):
                        kind = 9
                    elif not prompt and next_row and not next_row[2] and next_row[3] and not any(next_row[5:]):
                        kind = 2
                    else:
                        kind = 3 if prompt else 1
                else:
                    kind = 7
            else:
                try:
                    kind = int(declared)
                except ValueError as error:
                    raise ValueError(f'B{row}: specify a type from 1 to 9') from error
            if kind not in range(1,10):
                raise ValueError(f'B{row}: unknown type')
            if kind in (1,2,3,7,9):
                content = unit(stimulus,row,3)
                options, expected, labels = choices(raw,row)
                interaction = {'stimulus':content,'options':options}
                question = unit(prompt,row,4) if prompt else None
                if question:
                    interaction['question'] = question
                shown = (question or content).get('text','')
                if kind == 2:
                    header = (2,interaction,{label:key for key,label in labels.items()})
                    if index+1 == len(rows) or rows[index+1][1][2]:
                        raise ValueError(f'C{row}: matrix header needs subquestions')
                    continue
                if kind == 3:
                    header = (3, {'stimulus':content},None)
                add(row,kind,shown,interaction,{'mode':'choice','expected':expected,'labels':labels})
            elif kind in (4,5):
                segments, expected, labels = [], {}, {}
                position = 0
                matches = list(BRACKET.finditer(stimulus))
                if not matches or BRACKET.sub('',stimulus).count('[') or BRACKET.sub('',stimulus).count(']'):
                    raise ValueError(f'C{row}: malformed cloze brackets')
                for slot,match in enumerate(matches):
                    segments.append({'text':stimulus[position:match.start()]})
                    body=match.group(1)
                    typed=body.startswith('~')
                    if typed and kind == 4:
                        raise ValueError(f'C{row}: typed blanks require type 5')
                    candidates=[v.strip() for v in body.lstrip('~').split(';')]
                    if len(candidates)<2 or any(not v for v in candidates):
                        raise ValueError(f'C{row}: each blank needs at least two nonempty options')
                    key=str(slot)
                    if typed:
                        segments.append({'id':key,'typed':True})
                        expected[key]=candidates[0]
                        labels[key]={candidates[0]:candidates[0]}
                    else:
                        opts, correct, mapping=choices(candidates,row,3+slot)
                        segments.append({'id':key,'typed':False,'options':opts})
                        expected[key]=correct
                        labels[key]=mapping
                    position=match.end()
                segments.append({'text':stimulus[position:]})
                passage=unit(stimulus,row,3)
                add(row,kind,prompt,{'segments':segments, **({'speechUrl':passage['speechUrl']} if 'speechUrl' in passage else {})},{'mode':'cloze','expected':expected,'labels':labels,'typed':[s['id'] for s in segments if s.get('typed')]})
            elif kind == 6:
                tuples=[[v.strip() for v in value.split('*')] for value in raw]
                lengths={len(v) for v in tuples}
                if len(raw)<2 or len(lengths)!=1 or not lengths <= {2,3} or any(not v for group in tuples for v in group):
                    raise ValueError(f'E{row}: matching needs two or more pairs/triplets')
                if any(len({normalize(group[column]) for group in tuples}) != len(tuples) for column in range(len(tuples[0]))):
                    raise ValueError(f'E{row}: matching columns must contain distinct values')
                columns=[[] for _ in range(len(tuples[0]))]
                expected, labels={},{}
                for group,parts in enumerate(tuples):
                    for column,value in enumerate(parts):
                        token=secrets.token_hex(12)
                        columns[column].append({'id':token,'content':unit(value,row,5+group)})
                        labels[token]=value
                        if column:
                            expected[f'{columns[0][-1]["id"]}:{column}']=token
                for column in columns:
                    rng.shuffle(column)
                add(row,6,prompt,{'columns':columns},{'mode':'matching','expected':expected,'labels':labels})
            elif kind == 8:
                entries=[v.split('*',1) for v in raw]
                if len(entries)<2 or any(len(v)!=2 or not all(p.strip() for p in v) for v in entries):
                    raise ValueError(f'E{row}: media set needs label * media entries')
                options=[]
                for j,(label,media) in enumerate(entries):
                    content=unit(media.strip(),row,5+j)
                    if content['kind'] != 'media':
                        raise ValueError(f'{get_column_letter(5+j)}{row}: persistent set requires media')
                    options.append({'id':secrets.token_hex(12),'content':content})
                labels={o['id']:entries[j][0].strip() for j,o in enumerate(options)}
                for j,option in enumerate(options):
                    shuffled=list(options)
                    rng.shuffle(shuffled)
                    prompt_unit=unit(entries[j][0].strip(),row,5+j)
                    add(row,8,prompt_unit.get('text',''),{'question':prompt_unit,'options':shuffled,'persistent':True},{'mode':'choice','expected':option['id'],'labels':labels},str(j))
        except (ValueError, OSError) as error:
            warnings.append(f'{source}: {error}')
            header=None
    if not questions:
        raise ValueError('No valid Anvil questions. '+'; '.join(warnings[:20]))
    return {'engine':'anvil','schemaVersion':'1','items':items,'questions':questions,'warnings':warnings}


def score(question, answer):
    solution=question['solution']
    mode=solution['mode']
    expected=solution['expected']
    dimensions=[]
    submitted=[]
    answers=[]
    if mode == 'choice':
        correct=isinstance(answer,str) and answer == expected
        labels=solution['labels']
        submitted=[labels.get(answer,'—')]
        answers=[labels[expected]]
    else:
        try:
            actual=json.loads(answer)
        except (ValueError,TypeError):
            actual={}
        shape=isinstance(actual,dict) and set(actual)==set(expected) and all(isinstance(v,str) for v in actual.values())
        if not isinstance(actual,dict):
            actual={}
        for slot,value in expected.items():
            entered=actual.get(slot,'')
            if not isinstance(entered,str):
                entered=''
            typed=mode=='cloze' and slot in solution.get('typed',[])
            hit=normalize(entered)==normalize(value) if typed else entered==value
            mapping=solution['labels'].get(slot,{}) if mode=='cloze' else solution['labels']
            label=mapping.get(value,value)
            answers.append(label)
            submitted.append(entered if typed else mapping.get(entered,'—'))
            dimensions.append({'id':slot,'label':label,'score':int(hit)})
        correct=shape and all(d['score']==1 for d in dimensions)
    return {'score':int(correct),'fullyCorrect':bool(correct),'correctAnswers':answers,'matched':answers if correct else [],'missed':[] if correct else answers,'extra':[],'dimensions':dimensions,'submittedLabels':submitted}
