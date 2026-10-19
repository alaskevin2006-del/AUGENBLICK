"""Companion script compiler and completion-only practice evidence."""
import base64
import copy
import json
import math
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from .companion_audio import MAX_BYTES, features, compare, reference_bytes
from .common import safe_file

AUDIO_COLS = {'E': 'guidance-native', 'G': 'guidance-target', 'J': 'male1', 'L': 'female1', 'N': 'male2', 'P': 'female2'}
COLORS = {'red', 'green', 'blue', 'yellow', 'orange', 'purple'}


def display(raw):
    if raw == 'WIPE CLEAN':
        return {'operation': 'clear', 'spans': []}
    operation = 'append' if raw.startswith('*') else 'newline' if raw else 'none'
    text = raw[1:] if operation == 'append' else raw
    spans, style = [], {}
    for part in re.split(r'(\[/?(?:bold|red|green|blue|yellow|orange|purple)\])', text):
        if re.fullmatch(r'\[/?(?:bold|red|green|blue|yellow|orange|purple)\]', part):
            tag = part.strip('[]/')
            key = 'bold' if tag == 'bold' else 'color'
            if part.startswith('[/'):
                style.pop(key, None)
            else:
                style[key] = True if key == 'bold' else tag
        elif part:
            spans.append({'text': part, **style})
    return {'operation': operation, 'spans': spans}


def compile_source(root, manifest, asset):
    root = Path(root).resolve()
    source = safe_file(root, manifest.get('source', 'script.xlsx'))
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Companion requires .xlsx; convert legacy .xls first.')
    try:
        with ZipFile(source) as archive:
            if sum(member.file_size for member in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError('Workbook expanded content exceeds 64 MiB.')
        book = load_workbook(source, data_only=True, read_only=True)
    except (BadZipFile, OSError, KeyError, ParseError, InvalidFileException) as error:
        raise ValueError(f'{source.name}: invalid XLSX workbook.') from error
    questions, items, checkpoints, transcript, seen = [], [], [], [], set()
    reference_cache = {}
    try:
        if (book.active.max_column or 0) > 17:
            raise ValueError('Companion script must use at most 17 columns (A–Q).')
        if (book.active.max_row or 0) > 2000:
            raise ValueError('Companion script exceeds 2000 rows.')
        for number, row in enumerate(book.active.iter_rows(values_only=True), 1):
            if number > 2000:
                raise ValueError('Companion script exceeds 2000 rows.')
            if len(row) > 17:
                raise ValueError('Companion script must use at most 17 columns (A–Q).')
            cells = list(row) + [None] * 17
            if not any(v is not None for v in cells[1:17]):
                continue
            # Header is explicit; raw legacy scripts remain row-1 based.
            if number == 1 and str(cells[1]).lower() == 'checkpoint':
                continue
            try:
                checkpoint = cells[1]
                if checkpoint is not None:
                    if isinstance(checkpoint, bool) or float(checkpoint) != int(checkpoint) or int(checkpoint) < 1 or int(checkpoint) in seen:
                        raise ValueError('Checkpoint must be a unique positive integer.')
                    checkpoint = int(checkpoint)
                    seen.add(checkpoint)
                pause = float(cells[2] or 0)
                if not math.isfinite(pause) or not 0 <= pause <= 300:
                    raise ValueError('Pause must be 0–300 seconds.')
                voice = str(cells[7] or '').upper()
                if voice not in ('', 'Y'):
                    raise ValueError('H must be Y or empty.')
                refs = [(col, str(cells[ord(col) - 65]).strip()) for col in AUDIO_COLS if cells[ord(col) - 65]]
                if len(refs) > 1 or (voice and not refs):
                    raise ValueError('Use exactly one reference audio for voice input; audio columns are mutually exclusive.')
                ref, analysis, duration, speaker, reference_text = None, None, None, None, ''
                if refs:
                    col, name = refs[0]
                    reference_text = str(cells[ord(col) - 66] or '')
                    if not re.fullmatch(col + r'\d{4}[^/\\]*', name):
                        raise ValueError('Audio name must begin with its column letter and four digits.')
                    # Prefix identifiers allow descriptive suffixes, but ambiguity blocks publication.
                    candidates = [p for p in source.parent.iterdir() if p.is_file() and p.name[:5] == name[:5] and p.suffix.lower() in ('.wav', '.mp3', '.ogg', '.m4a')]
                    if len(candidates) != 1:
                        raise ValueError('Reference audio is missing or ambiguous: ' + name[:5])
                    path = safe_file(root, candidates[0].relative_to(root).as_posix())
                    if path not in reference_cache:
                        reference_cache[path] = features(reference_bytes(path))
                    analysis = reference_cache[path]
                    duration = round(analysis['duration'] * 1000)
                    if voice and duration > 20000:
                        raise ValueError('Voice reference must be at most 20 seconds.')
                    ref, speaker = asset(path.relative_to(root).as_posix()), AUDIO_COLS[col]
                row_id = f'companion-row-{number}'
                instruction = display(str(cells[16] or ''))
                interaction = {'voiceInput': bool(voice), 'pauseSeconds': pause, 'reference': ref, 'referenceDurationMs': duration, 'speaker': speaker, 'checkpoint': checkpoint, 'displayBefore': copy.deepcopy(transcript), 'display': instruction, 'assessment': 'completion-only'}
                if instruction['operation'] == 'clear':
                    transcript = []
                elif instruction['operation'] in ('append', 'newline'):
                    segments = [{**span, 'audio': ref} for span in instruction['spans']]
                    if instruction['operation'] == 'append' and transcript:
                        transcript[-1].extend(segments)
                    else:
                        transcript.append(segments)
                interaction['displayAfter'] = copy.deepcopy(transcript)
                name = ''.join(span['text'] for span in instruction['spans']) or f'#{number}'
                items.append({'id': row_id, 'name': name[:255], 'values': {}})
                questions.append({'objectId': row_id, 'kind': 'companion-row', 'prompt': name, 'interaction': interaction, 'solution': {'features': analysis, 'referenceText': reference_text}})
                if checkpoint is not None:
                    checkpoints.append({'id': checkpoint, 'objectId': row_id, 'label': str(checkpoint)})
            except (ValueError, TypeError, OverflowError) as error:
                raise ValueError(f'{source.name} row {number}: {error}') from error
    except (BadZipFile, OSError, KeyError, ParseError, InvalidFileException) as error:
        raise ValueError(f'{source.name}: invalid XLSX workbook.') from error
    finally:
        book.close()
    if not questions:
        raise ValueError('Companion script has no executable rows.')
    return {'engine': 'companion', 'items': items, 'questions': questions, 'checkpoints': checkpoints, 'ordered': True, 'assessment': 'completion-only'}


def score(question, answer):
    if len(answer) > 2800000:
        raise ValueError('Recording exceeds 2 MiB.')
    try:
        submitted = json.loads(answer)
    except (ValueError, TypeError) as error:
        raise ValueError('Invalid Companion response.') from error
    if not isinstance(submitted, dict):
        raise ValueError('Invalid Companion response.')
    diagnostic, status = None, 'completed'
    if question['interaction']['voiceInput']:
        if submitted.get('noResponse') is True:
            status = 'no-response'
        else:
            try:
                raw = base64.b64decode(submitted['audio'], validate=True)
            except (KeyError, ValueError, TypeError) as error:
                raise ValueError('A PCM WAV recording is required.') from error
            if len(raw) > MAX_BYTES:
                raise ValueError('Recording exceeds 2 MiB.')
            actual = features(raw)
            reference = question['solution']['features']
            if actual['duration'] > reference['duration'] * 1.5 + .25:
                raise ValueError('Recording exceeds the allowed reference duration.')
            diagnostic = compare(reference, actual)
            status = 'recorded' if actual['speechFraction'] >= .05 else 'no-response'
    elif submitted.get('ack') is not True:
        raise ValueError('Acknowledge the completed script row.')
    result = {'score': 1, 'fullyCorrect': True, 'correctAnswers': [], 'matched': [], 'missed': [], 'extra': [], 'dimensions': [], 'submittedLabels': ['[recording]' if status == 'recorded' else status], 'assessment': 'completion-only', 'responseStatus': status}
    if diagnostic is not None:
        result['diagnostics'] = diagnostic
    return result
